"""
Mi Negocio - Flet + Supabase (Con Ventana de Descarga de Excel)

Versiones probadas: flet==0.28.3, supabase>=2.9, openpyxl>=3.1
Ejecutar en PC:  pip install "flet==0.28.3" "supabase>=2.9" openpyxl && python mi_negocio_flet.py
Empaquetar APK:  en pyproject.toml -> dependencies = ["flet==0.28.3", "supabase", "openpyxl"]
"""
import asyncio
from datetime import date

import flet as ft
from supabase import AsyncClient, acreate_client

# Librerías para generación de Excel estilizado
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ====================== CREDENCIALES SUPABASE ======================
SUPABASE_URL = "https://pgzpuysruhogvhiinzor.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InBnenB1eXNydWhvZ3ZoaWluem9yIiwicm9sZSI6ImFub24iLCJpYXQiOjE3OTE2NjU3NDgsImV4cCI6MjEwNzI0MTc0OH0.IWBb2bNS-c7YRlB6jysqA7rC_ceGA2AQbzf_RQ0wugM"
# ===================================================================

MONEDA = "S/"
TABLAS = ("productos", "ventas", "gastos", "config")
VERDE = ft.Colors.TEAL_700
LIMITE_LISTA = 100


def money(x: float) -> str:
    return f"{MONEDA} {x:,.2f}"


def num(texto, defecto=None):
    try:
        return float(str(texto).replace(",", ".").strip())
    except (ValueError, AttributeError):
        return defecto


async def traer_todo(client: AsyncClient, tabla: str, columnas="*", orden="id"):
    """Trae todas las filas de la tabla."""
    filas, desde, paso = [], 0, 1000
    while True:
        r = (
            await client.table(tabla)
            .select(columnas)
            .order(orden, desc=True)
            .range(desde, desde + paso - 1)
            .execute()
        )
        filas.extend(r.data)
        if len(r.data) < paso:
            return filas
        desde += paso


async def main(page: ft.Page):
    page.title = "Mi Negocio"
    page.theme = ft.Theme(color_scheme_seed=ft.Colors.TEAL)
    page.theme_mode = ft.ThemeMode.LIGHT
    page.padding = 12
    page.bgcolor = ft.Colors.BLUE_GREY_50
    page.horizontal_alignment = ft.CrossAxisAlignment.CENTER

    if "TU-PROYECTO" in SUPABASE_URL or "TU-ANON" in SUPABASE_KEY:
        page.add(
            ft.SafeArea(
                ft.Container(
                    padding=24,
                    content=ft.Text(
                        "⚠️ Pega tu SUPABASE_URL y SUPABASE_KEY al inicio del archivo.",
                        size=20,
                    ),
                )
            )
        )
        return

    client: AsyncClient = await acreate_client(SUPABASE_URL, SUPABASE_KEY)

    # ------------------------------------------------------------------ estado
    estado = {"capital": 0.0, "productos": [], "ventas": [], "gastos": []}
    lock = asyncio.Lock()

    def aviso(texto, error=False):
        page.open(
            ft.SnackBar(
                ft.Text(texto, color=ft.Colors.WHITE),
                bgcolor=ft.Colors.RED_700 if error else ft.Colors.GREEN_700,
            )
        )

    def abrir_calendario(campo_destino):
        """Abre un calendario flotante para seleccionar fecha fácilmente."""
        def al_seleccionar_fecha(e):
            if e.control.value:
                campo_destino.value = e.control.value.strftime("%Y-%m-%d")
                page.update()

        page.open(
            ft.DatePicker(
                first_date=date(2024, 1, 1),
                last_date=date(2030, 12, 31),
                on_change=al_seleccionar_fecha,
            )
        )

    # ---------------------------------------------------------------- resumen
    def resumen():
        v = estado["ventas"]
        vendido = sum((x["cantidad"] or 0) * (x["precio"] or 0) for x in v)
        cobrado = sum((x["cantidad"] or 0) * (x["precio"] or 0) for x in v if x["pagado"] == 1)
        gastos = sum(x["monto"] or 0 for x in estado["gastos"])
        capital = estado["capital"]
        return dict(
            capital=capital,
            vendido=vendido,
            cobrado=cobrado,
            por_cobrar=vendido - cobrado,
            gastos=gastos,
            ganancia=cobrado - gastos,
            caja=capital + cobrado - gastos,
        )

    def tarjeta(titulo, valor, color, grande=False):
        return ft.Card(
            color=color,
            margin=0,
            content=ft.Container(
                padding=12,
                content=ft.Column(
                    [
                        ft.Text(titulo, size=13, color=ft.Colors.WHITE70, overflow=ft.TextOverflow.ELLIPSIS),
                        ft.Text(valor, size=26 if grande else 20, weight=ft.FontWeight.BOLD, color=ft.Colors.WHITE, overflow=ft.TextOverflow.ELLIPSIS),
                    ],
                    spacing=2,
                ),
            ),
        )

    tarjetas = ft.Column(spacing=6)

    def pintar_tarjetas():
        r = resumen()
        tarjetas.controls = [
            ft.Row([ft.Container(tarjeta("💰 Dinero en caja", money(r["caja"]), ft.Colors.TEAL_800, True), expand=True)]),
            ft.Row([
                ft.Container(tarjeta("🕒 Me deben", money(r["por_cobrar"]), ft.Colors.ORANGE_600), expand=True),
                ft.Container(tarjeta("🛒 Vendido", money(r["vendido"]), ft.Colors.LIGHT_BLUE_700), expand=True),
            ], spacing=6),
            ft.Row([
                ft.Container(tarjeta("💸 Gastos", money(r["gastos"]), ft.Colors.RED_600), expand=True),
                ft.Container(tarjeta("📈 Ganancia", money(r["ganancia"]), ft.Colors.GREEN_700), expand=True),
            ], spacing=6),
        ]

    # --------------------------------------------------------- confirmar borrado
    def confirmar_borrado(tabla, id_, texto):
        async def si(e):
            page.close(dlg)
            try:
                await client.table(tabla).delete().eq("id", id_).execute()
                await refrescar()
            except Exception as ex:
                aviso(f"Error: {ex}", True)

        def no(e):
            page.close(dlg)

        dlg = ft.AlertDialog(
            modal=True,
            title=ft.Text("¿Eliminar?"),
            content=ft.Text(texto),
            actions=[ft.TextButton("Cancelar", on_click=no), ft.FilledButton("Eliminar", on_click=si)],
        )
        page.open(dlg)

    def campo(label, **kw):
        return ft.TextField(
            label=label,
            border_radius=12,
            filled=True,
            bgcolor=ft.Colors.WHITE,
            border_color=ft.Colors.GREY_300,
            focused_border_color=ft.Colors.TEAL_700,
            text_size=15,
            content_padding=12,
            **kw
        )

    def boton(texto, icono, accion, color=ft.Colors.TEAL_800):
        return ft.Row(
            [
                ft.FilledButton(
                    texto,
                    icon=icono,
                    on_click=accion,
                    height=48,
                    style=ft.ButtonStyle(
                        shape=ft.RoundedRectangleBorder(radius=12),
                        bgcolor=color,
                    ),
                    expand=True,
                )
            ]
        )

    # ======================================================================
    #                 GENERACIÓN DE REPORTE CON FILEPICKER
    # ======================================================================
    def guardar_excel_archivo(e: ft.FilePickerResultEvent):
        """Callback ejecutado cuando el usuario selecciona la ruta de guardado."""
        if not e.path:
            return
        
        ruta_destino = e.path
        if not ruta_destino.endswith(".xlsx"):
            ruta_destino += ".xlsx"

        try:
            wb = openpyxl.Workbook()
            
            # Estilos visuales
            HEADER_FILL = PatternFill(start_color="004D40", end_color="004D40", fill_type="solid")
            HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="004D40")
            BOLD_FONT = Font(name="Calibri", size=11, bold=True)
            THIN_BORDER = Border(
                left=Side(style='thin', color='DDDDDD'),
                right=Side(style='thin', color='DDDDDD'),
                top=Side(style='thin', color='DDDDDD'),
                bottom=Side(style='thin', color='DDDDDD')
            )
            FORMATO_MONEDA = '"S/" #,##0.00'

            # ---------------- 1. Pestaña Resumen
            ws_resumen = wb.active
            ws_resumen.title = "Resumen General"
            ws_resumen.views.sheetView[0].showGridLines = True
            
            ws_resumen.append(["REPORTE GENERAL DEL NEGOCIO"])
            ws_resumen.cell(row=1, column=1).font = TITLE_FONT
            ws_resumen.append([f"Fecha de emisión: {date.today()}"])
            ws_resumen.append([])
            
            ws_resumen.append(["Indicador Financiero", "Monto"])
            r = resumen()
            indicadores = [
                ("Capital Inicial", r["capital"]),
                ("Vendido Total", r["vendido"]),
                ("Cobrado", r["cobrado"]),
                ("Por Cobrar (Pendiente)", r["por_cobrar"]),
                ("Gastos Totales", r["gastos"]),
                ("Ganancia Neta", r["ganancia"]),
                ("Dinero en Caja Actual", r["caja"]),
            ]
            for ind, val in indicadores:
                ws_resumen.append([ind, val])
                
            for col in range(1, 3):
                cell = ws_resumen.cell(row=4, column=col)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
            for row in range(5, 5 + len(indicadores)):
                c1 = ws_resumen.cell(row=row, column=1)
                c2 = ws_resumen.cell(row=row, column=2)
                c1.border = THIN_BORDER
                c2.border = THIN_BORDER
                c2.number_format = FORMATO_MONEDA
                if row == 5 + len(indicadores) - 1:
                    c1.font = BOLD_FONT
                    c2.font = BOLD_FONT

            # ---------------- 2. Pestaña Ventas
            ws_ventas = wb.create_sheet(title="Ventas")
            ws_ventas.views.sheetView[0].showGridLines = True
            ws_ventas.append(["Fecha", "Cliente", "Producto", "Cantidad", "Precio Unitario", "Total", "Estado"])
            for col in range(1, 8):
                cell = ws_ventas.cell(row=1, column=col)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
            for v in estado["ventas"]:
                cant = v["cantidad"] or 0
                prec = v["precio"] or 0
                total = cant * prec
                estado_txt = "PAGADO" if v["pagado"] == 1 else "PENDIENTE"
                ws_ventas.append([v["fecha"], v["cliente"], v["producto"], cant, prec, total, estado_txt])
                
            for row in range(2, ws_ventas.max_row + 1):
                for col in range(1, 8):
                    c = ws_ventas.cell(row=row, column=col)
                    c.border = THIN_BORDER
                    if col in (5, 6):
                        c.number_format = FORMATO_MONEDA
                    if col == 7:
                        c.alignment = Alignment(horizontal="center")
                        c.font = Font(color="008000" if c.value == "PAGADO" else "D97706", bold=True)

            # ---------------- 3. Pestaña Gastos
            ws_gastos = wb.create_sheet(title="Gastos")
            ws_gastos.views.sheetView[0].showGridLines = True
            ws_gastos.append(["Fecha", "Concepto", "Monto"])
            for col in range(1, 4):
                cell = ws_gastos.cell(row=1, column=col)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
            for g in estado["gastos"]:
                ws_gastos.append([g["fecha"], g["concepto"], g["monto"] or 0])
                
            for row in range(2, ws_gastos.max_row + 1):
                for col in range(1, 4):
                    c = ws_gastos.cell(row=row, column=col)
                    c.border = THIN_BORDER
                    if col == 3:
                        c.number_format = FORMATO_MONEDA

            # ---------------- 4. Pestaña Productos
            ws_prods = wb.create_sheet(title="Productos")
            ws_prods.views.sheetView[0].showGridLines = True
            ws_prods.append(["Nombre", "Precio Venta", "Costo", "Ganancia x Unid."])
            for col in range(1, 5):
                cell = ws_prods.cell(row=1, column=col)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
            for p in estado["productos"]:
                pv = p["precio"] or 0
                pc = p["costo"] or 0
                ws_prods.append([p["nombre"], pv, pc, pv - pc])
                
            for row in range(2, ws_prods.max_row + 1):
                for col in range(1, 5):
                    c = ws_prods.cell(row=row, column=col)
                    c.border = THIN_BORDER
                    if col in (2, 3, 4):
                        c.number_format = FORMATO_MONEDA

            # Auto-ajuste de anchos de columna
            for ws in [ws_resumen, ws_ventas, ws_gastos, ws_prods]:
                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(ruta_destino)
            aviso("¡Reporte en Excel generado con éxito! 📊")

        except Exception as ex:
            aviso(f"Error al generar Excel: {ex}", True)

    # Componente de FilePicker para solicitar ubicación de guardado
    file_picker = ft.FilePicker(on_result=guardar_excel_archivo)
    page.overlay.append(file_picker)

    def solicitar_guardar_excel(e=None):
        nombre_defecto = f"Reporte_Mi_Negocio_{date.today().strftime('%Y%m%d')}.xlsx"
        file_picker.save_file(
            file_name=nombre_defecto,
            allowed_extensions=["xlsx"],
            dialog_title="Guardar Reporte Excel",
        )

    # ======================================================================
    #                               VENTAS
    # ======================================================================
    def al_cambiar_producto(e=None):
        if v_producto.value:
            p = next((x for x in estado["productos"] if x["nombre"] == v_producto.value), None)
            if p and p["precio"] is not None:
                v_precio.value = f"{p['precio']:g}"
            else:
                v_precio.value = ""
        else:
            v_precio.value = ""
        page.update()

    v_fecha = campo("Fecha (AAAA-MM-DD)", value=str(date.today()), keyboard_type=ft.KeyboardType.DATETIME, expand=True)
    btn_cal_ventas = ft.IconButton(
        icon=ft.Icons.CALENDAR_MONTH,
        icon_color=ft.Colors.TEAL_800,
        tooltip="Elegir fecha",
        on_click=lambda _: abrir_calendario(v_fecha),
    )

    v_cliente = campo("Nombre del cliente", capitalization=ft.TextCapitalization.WORDS)
    v_producto = ft.Dropdown(
        label="Producto",
        hint_text="Seleccionar producto...",
        border_radius=12,
        filled=True,
        bgcolor=ft.Colors.WHITE,
        border_color=ft.Colors.GREY_300,
        focused_border_color=ft.Colors.TEAL_700,
        text_size=15,
        content_padding=12,
        options=[],
        on_change=al_cambiar_producto,
        expand=True,
    )
    v_cantidad = campo("Cantidad", value="1", keyboard_type=ft.KeyboardType.NUMBER, expand=True)
    v_precio = campo("Precio de venta", keyboard_type=ft.KeyboardType.NUMBER, expand=True)
    v_pagado = ft.Switch(value=False, active_color=ft.Colors.TEAL_700)

    v_pagado_container = ft.Container(
        content=ft.Row(
            [
                ft.Text("¿Ya pagó el cliente?", size=15, color=ft.Colors.GREY_800),
                v_pagado,
            ],
            alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
        ),
        padding=ft.padding.symmetric(horizontal=4),
    )

    lista_ventas = ft.Column(spacing=0)

    async def guardar_venta(e):
        cant = num(v_cantidad.value)
        if not v_producto.value or not cant or cant <= 0:
            aviso("Elige un producto y la cantidad", True)
            return
        precio = num(v_precio.value)
        if precio is None:
            p = next((x for x in estado["productos"] if x["nombre"] == v_producto.value), None)
            precio = p["precio"] if p else 0
        try:
            await client.table("ventas").insert(
                {
                    "fecha": v_fecha.value or str(date.today()),
                    "cliente": (v_cliente.value or "").strip() or "Sin nombre",
                    "producto": v_producto.value,
                    "cantidad": cant,
                    "precio": precio,
                    "pagado": 1 if v_pagado.value else 0,
                }
            ).execute()
        except Exception as ex:
            aviso(f"Error: {ex}", True)
            return

        v_cliente.value = ""
        v_producto.value = None
        v_cantidad.value = "1"
        v_precio.value = ""
        v_pagado.value = False

        aviso("Venta guardada ✅")
        await refrescar()

    def pintar_ventas():
        if not estado["ventas"]:
            lista_ventas.controls = [ft.Text("Aún no hay ventas. Registra la primera arriba 👆", color=ft.Colors.GREY)]
            return
        items = []
        for v in estado["ventas"][:LIMITE_LISTA]:
            total = (v["cantidad"] or 0) * (v["precio"] or 0)
            pagado = v["pagado"] == 1

            async def alternar(e, id_=v["id"]):
                try:
                    await client.table("ventas").update({"pagado": 1 if e.control.value else 0}).eq("id", id_).execute()
                    await refrescar()
                except Exception as ex:
                    aviso(f"Error: {ex}", True)

            items.append(
                ft.Card(
                    content=ft.Container(
                        border=ft.border.only(left=ft.BorderSide(6, ft.Colors.GREEN if pagado else ft.Colors.ORANGE)),
                        border_radius=12,
                        padding=ft.padding.only(left=8, right=8, top=4, bottom=4),
                        content=ft.Column(
                            [
                                ft.ListTile(
                                    content_padding=0,
                                    leading=ft.Icon(
                                        ft.Icons.CHECK_CIRCLE if pagado else ft.Icons.SCHEDULE,
                                        color=ft.Colors.GREEN if pagado else ft.Colors.ORANGE,
                                        size=28,
                                    ),
                                    title=ft.Text(
                                        f"{v['cliente']} · {v['producto']}",
                                        weight=ft.FontWeight.W_500,
                                        size=15,
                                        overflow=ft.TextOverflow.ELLIPSIS,
                                    ),
                                    subtitle=ft.Text(
                                        f"{v['fecha']} · {v['cantidad']:g} × {money(v['precio'] or 0)}",
                                        overflow=ft.TextOverflow.ELLIPSIS,
                                    ),
                                    trailing=ft.Text(money(total), size=16, weight=ft.FontWeight.BOLD),
                                ),
                                ft.Row(
                                    [
                                        ft.Text("Pagó", size=14, color=ft.Colors.GREY_700),
                                        ft.Switch(value=pagado, on_change=alternar, active_color=ft.Colors.GREEN),
                                        ft.IconButton(
                                            ft.Icons.DELETE_OUTLINE,
                                            icon_size=24,
                                            on_click=lambda e, i=v["id"], t=f"{v['cliente']} · {money(total)}":
                                                confirmar_borrado("ventas", i, t),
                                        ),
                                    ],
                                    alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
                                ),
                            ],
                            spacing=0,
                        ),
                    )
                )
            )
        lista_ventas.controls = items

    tab_ventas = ft.Column(
        [
            ft.Card(
                content=ft.Container(
                    padding=14,
                    content=ft.Column(
                        [
                            ft.Text("Nueva venta", size=18, weight=ft.FontWeight.BOLD),
                            ft.Row([v_fecha, btn_cal_ventas], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
                            v_cliente,
                            ft.Row([v_producto]),
                            ft.Row([v_cantidad, v_precio], spacing=8),
                            v_pagado_container,
                            boton("Guardar venta", ft.Icons.ADD, guardar_venta),
                        ],
                        spacing=10,
                    ),
                )
            ),
            lista_ventas,
        ],
        scroll=ft.ScrollMode.AUTO, expand=True, spacing=8,
    )

    # ======================================================================
    #                               GASTOS
    # ======================================================================
    g_fecha = campo("Fecha (AAAA-MM-DD)", value=str(date.today()), keyboard_type=ft.KeyboardType.DATETIME, expand=True)
    btn_cal_gastos = ft.IconButton(
        icon=ft.Icons.CALENDAR_MONTH,
        icon_color=ft.Colors.TEAL_800,
        tooltip="Elegir fecha",
        on_click=lambda _: abrir_calendario(g_fecha),
    )
    g_concepto = campo("¿En qué gastaste?", capitalization=ft.TextCapitalization.SENTENCES)
    g_monto = campo("Monto", keyboard_type=ft.KeyboardType.NUMBER)
    lista_gastos = ft.Column(spacing=0)

    async def guardar_gasto(e):
        monto = num(g_monto.value)
        if not g_concepto.value or not monto or monto <= 0:
            aviso("Escribe el concepto y el monto", True)
            return
        try:
            await client.table("gastos").insert(
                {"fecha": g_fecha.value or str(date.today()), "concepto": g_concepto.value.strip(), "monto": monto}
            ).execute()
        except Exception as ex:
            aviso(f"Error: {ex}", True)
            return
        g_concepto.value, g_monto.value = "", ""
        aviso("Gasto guardado ✅")
        await refrescar()

    def pintar_gastos():
        if not estado["gastos"]:
            lista_gastos.controls = [ft.Text("Aún no hay gastos.", color=ft.Colors.GREY)]
            return
        lista_gastos.controls = [
            ft.Card(
                content=ft.ListTile(
                    leading=ft.Icon(ft.Icons.PAYMENTS_OUTLINED, color=ft.Colors.RED_400, size=28),
                    title=ft.Text(g["concepto"], weight=ft.FontWeight.W_500, size=15, overflow=ft.TextOverflow.ELLIPSIS),
                    subtitle=ft.Text(g["fecha"]),
                    trailing=ft.Row(
                        [
                            ft.Text(money(g["monto"] or 0), size=16, weight=ft.FontWeight.BOLD, color=ft.Colors.RED_600),
                            ft.IconButton(
                                ft.Icons.DELETE_OUTLINE, icon_size=24,
                                on_click=lambda e, i=g["id"], t=f"{g['concepto']} · {money(g['monto'] or 0)}":
                                    confirmar_borrado("gastos", i, t),
                            ),
                        ],
                        tight=True,
                    ),
                )
            )
            for g in estado["gastos"][:LIMITE_LISTA]
        ]

    tab_gastos = ft.Column(
        [
            ft.Card(
                content=ft.Container(
                    padding=14,
                    content=ft.Column(
                        [
                            ft.Text("Nuevo gasto", size=18, weight=ft.FontWeight.BOLD),
                            ft.Row([g_fecha, btn_cal_gastos], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
                            g_concepto, g_monto,
                            boton("Guardar gasto", ft.Icons.ADD, guardar_gasto),
                        ],
                        spacing=10,
                    ),
                )
            ),
            lista_gastos,
        ],
        scroll=ft.ScrollMode.AUTO, expand=True, spacing=8,
    )

    # ======================================================================
    #                             PRODUCTOS
    # ======================================================================
    p_nombre = campo("Nombre", capitalization=ft.TextCapitalization.WORDS)
    p_precio = campo("Precio de venta", keyboard_type=ft.KeyboardType.NUMBER, expand=True)
    p_costo = campo("Lo que me cuesta", keyboard_type=ft.KeyboardType.NUMBER, expand=True)
    lista_productos = ft.Column(spacing=0)

    async def guardar_producto(e):
        nombre, precio = (p_nombre.value or "").strip(), num(p_precio.value)
        if not nombre or precio is None:
            aviso("Escribe el nombre y el precio", True)
            return
        try:
            await client.table("productos").upsert(
                {"nombre": nombre, "precio": precio, "costo": num(p_costo.value, 0)}, on_conflict="nombre"
            ).execute()
        except Exception as ex:
            aviso(f"Error: {ex}", True)
            return
        p_nombre.value, p_precio.value, p_costo.value = "", "", ""
        aviso("Producto guardado ✅")
        await refrescar()

    def pintar_productos():
        if not estado["productos"]:
            lista_productos.controls = [ft.Text("Agrega tus productos arriba para empezar 👆", color=ft.Colors.GREY)]
            return

        def editar(p):
            p_nombre.value, p_precio.value, p_costo.value = p["nombre"], f"{p['precio']:g}", f"{(p['costo'] or 0):g}"
            page.update()

        lista_productos.controls = [
            ft.Card(
                content=ft.ListTile(
                    leading=ft.Icon(ft.Icons.INVENTORY_2_OUTLINED, color=VERDE, size=28),
                    title=ft.Text(p["nombre"], weight=ft.FontWeight.W_500, size=15, overflow=ft.TextOverflow.ELLIPSIS),
                    subtitle=ft.Text(
                        f"Precio {money(p['precio'] or 0)} · Me cuesta {money(p['costo'] or 0)}\n"
                        f"Gano {money((p['precio'] or 0) - (p['costo'] or 0))}",
                        overflow=ft.TextOverflow.ELLIPSIS,
                    ),
                    is_three_line=True,
                    on_click=lambda e, p=p: editar(p),
                    trailing=ft.IconButton(
                        ft.Icons.DELETE_OUTLINE, icon_size=24,
                        on_click=lambda e, i=p["id"], t=p["nombre"]: confirmar_borrado("productos", i, t),
                    ),
                )
            )
            for p in sorted(estado["productos"], key=lambda x: x["nombre"].lower())
        ]

    tab_productos = ft.Column(
        [
            ft.Card(
                content=ft.Container(
                    padding=14,
                    content=ft.Column(
                        [
                            ft.Text("Producto (si ya existe, se actualiza)", size=18, weight=ft.FontWeight.BOLD),
                            p_nombre,
                            ft.Row([p_precio, p_costo], spacing=8),
                            boton("Guardar producto", ft.Icons.ADD, guardar_producto),
                        ],
                        spacing=10,
                    ),
                )
            ),
            ft.Text("Toca un producto para editarlo", size=12, color=ft.Colors.GREY),
            lista_productos,
        ],
        scroll=ft.ScrollMode.AUTO, expand=True, spacing=8,
    )

    # ======================================================================
    #                               CAPITAL
    # ======================================================================
    c_capital = campo("Capital inicial", keyboard_type=ft.KeyboardType.NUMBER)

    async def guardar_capital(e):
        valor = num(c_capital.value)
        if valor is None:
            aviso("Escribe un número válido", True)
            return
        try:
            await client.table("config").upsert({"clave": "capital", "valor": valor}, on_conflict="clave").execute()
        except Exception as ex:
            aviso(f"Error: {ex}", True)
            return
        aviso("Capital actualizado ✅")
        await refrescar()

    tab_capital = ft.Column(
        [
            ft.Card(
                content=ft.Container(
                    padding=14,
                    content=ft.Column(
                        [
                            ft.Text("Configuración", size=18, weight=ft.FontWeight.BOLD),
                            c_capital,
                            boton("Guardar capital", ft.Icons.SAVE, guardar_capital),
                        ],
                        spacing=10,
                    ),
                )
            ),
            ft.Card(
                content=ft.Container(
                    padding=14,
                    content=ft.Column(
                        [
                            ft.Text("Reportes y Exportación", size=18, weight=ft.FontWeight.BOLD),
                            ft.Text("Descarga un reporte detallado en Excel con tus ventas, gastos y resumen.", size=13, color=ft.Colors.GREY_700),
                            boton("Exportar a Excel", ft.Icons.TABLE_CHART, solicitar_guardar_excel, color=ft.Colors.GREEN_800),
                        ],
                        spacing=10,
                    ),
                )
            ),
        ],
        scroll=ft.ScrollMode.AUTO, expand=True, spacing=8,
    )

    # ======================================================================
    #                      CARGA DE DATOS Y TIEMPO REAL
    # ======================================================================
    async def refrescar():
        async with lock:
            try:
                cfg = await client.table("config").select("valor").eq("clave", "capital").execute()
                estado["capital"] = float(cfg.data[0]["valor"]) if cfg.data else 0.0
                estado["productos"] = await traer_todo(client, "productos")
                estado["ventas"] = await traer_todo(client, "ventas")
                estado["gastos"] = await traer_todo(client, "gastos")
            except Exception as ex:
                aviso(f"Sin conexión: {ex}", True)
                return

            pintar_tarjetas()
            pintar_ventas()
            pintar_gastos()
            pintar_productos()

            nombres = [p["nombre"] for p in sorted(estado["productos"], key=lambda x: x["nombre"].lower())]
            v_producto.options = [ft.dropdown.Option(n) for n in nombres]

            if v_producto.value not in nombres:
                v_producto.value = None
                v_precio.value = ""

            if not c_capital.value:
                c_capital.value = f"{estado['capital']:g}"
            page.update()

    def hay_cambio(payload):
        page.run_task(refrescar)

    async def suscribir_tiempo_real():
        try:
            canal = client.channel("mi-negocio")
            for t in TABLAS:
                canal.on_postgres_changes("*", callback=hay_cambio, table=t, schema="public")
            await canal.subscribe()
        except Exception:
            pass

    async def sondeo():
        while True:
            await asyncio.sleep(20)
            await refrescar()

    # ------------------------------------------------------------------ inicio
    try:
        existe = await client.table("config").select("clave").eq("clave", "capital").execute()
        if not existe.data:
            await client.table("config").insert({"clave": "capital", "valor": 120}).execute()
    except Exception as ex:
        aviso(f"No se pudo conectar a Supabase: {ex}", True)

    tabs = ft.Tabs(
        selected_index=0,
        animation_duration=250,
        expand=True,
        tabs=[
            ft.Tab(text="Ventas", icon=ft.Icons.SHOPPING_CART, content=ft.Container(tab_ventas, padding=4)),
            ft.Tab(text="Gastos", icon=ft.Icons.PAYMENTS, content=ft.Container(tab_gastos, padding=4)),
            ft.Tab(text="Productos", icon=ft.Icons.INVENTORY_2, content=ft.Container(tab_productos, padding=4)),
            ft.Tab(text="Capital", icon=ft.Icons.SETTINGS, content=ft.Container(tab_capital, padding=4)),
        ],
    )

    page.add(
        ft.SafeArea(
            expand=True,
            content=ft.Container(
                width=480,
                content=ft.Column(
                    [
                        ft.Container(
                            padding=ft.padding.only(left=4, right=4, top=4),
                            content=ft.Text("🏪 Mi Negocio", size=24, weight=ft.FontWeight.BOLD, color=ft.Colors.TEAL_900),
                        ),
                        tarjetas,
                        tabs,
                    ],
                    expand=True,
                    spacing=6,
                ),
            ),
        )
    )

    await refrescar()
    await suscribir_tiempo_real()
    page.run_task(sondeo)


ft.app(target=main)
